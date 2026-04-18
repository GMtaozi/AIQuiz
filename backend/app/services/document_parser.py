"""Document Parser Service - 文档解析服务

支持解析 PDF、Word、Markdown、TXT 等格式的文档，
提取文本内容用于 AI 分析。
"""

import io
import re
import os
import tempfile
from typing import Optional

# 文档解析依赖（可选）
PDF_AVAILABLE = False
DOCX_AVAILABLE = False
WIN32_AVAILABLE = False

try:
    import win32com.client
    import pythoncom
    WIN32_AVAILABLE = True
except ImportError:
    pass

try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    pass

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    pass


def extract_text_from_pdf(file_content: bytes) -> str:
    """从 PDF 文件提取文本"""
    if not PDF_AVAILABLE:
        raise ImportError("PyPDF2 未安装，请运行: pip install PyPDF2")

    text_parts = []
    reader = PyPDF2.PdfReader(io.BytesIO(file_content))

    for page_num, page in enumerate(reader.pages):
        try:
            text = page.extract_text()
            if text:
                text_parts.append(f"[第{page_num + 1}页]\n{text}")
        except Exception:
            continue

    return "\n\n".join(text_parts)


def extract_text_from_docx(file_content: bytes) -> str:
    """从 Word .docx 文档提取文本"""
    if not DOCX_AVAILABLE:
        raise ImportError("python-docx 未安装，请运行: pip install python-docx")

    doc = docx.Document(io.BytesIO(file_content))
    paragraphs = []

    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    # 提取表格内容
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(f"[表格] {row_text}")

    return "\n\n".join(paragraphs)


def extract_text_from_doc(file_content: bytes) -> str:
    """从 Word .doc 文档（二进制格式）提取文本

    使用 Windows COM 接口通过 Microsoft Word 读取
    """
    if not WIN32_AVAILABLE:
        raise ImportError("pywin32 未安装，请运行: pip install pywin32")

    pythoncom.CoInitialize()

    try:
        # 保存到临时文件
        with tempfile.NamedTemporaryFile(suffix='.doc', delete=False) as tmp:
            tmp.write(file_content)
            tmp_path = tmp.name

        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            word.DisplayAlerts = False

            try:
                doc = word.Documents.Open(os.path.abspath(tmp_path))
                try:
                    # 提取纯文本
                    text = doc.Content.Text
                    # 清理文本
                    paragraphs = [p.strip() for p in text.split('\r') if p.strip()]
                    return '\n\n'.join(paragraphs)
                finally:
                    doc.Close(False)
            finally:
                word.Quit()
        finally:
            os.unlink(tmp_path)
    finally:
        pythoncom.CoUninitialize()


def extract_text_from_markdown(content: str) -> str:
    """从 Markdown 提取文本"""
    # 移除代码块
    content = re.sub(r'```[\s\S]*?```', '', content)
    # 移除行内代码
    content = re.sub(r'`[^`]+`', '', content)
    # 移除图片
    content = re.sub(r'!\[.*?\]\(.*?\)', '', content)
    # 移除链接，保留文本
    content = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', content)
    # 移除 HTML 标签
    content = re.sub(r'<[^>]+>', '', content)
    # 移除 Markdown 标题符号
    content = re.sub(r'^#{1,6}\s+', '', content, flags=re.MULTILINE)
    # 移除加粗和斜体
    content = re.sub(r'\*{1,3}([^*]+)\*{1,3}', r'\1', content)
    content = re.sub(r'_{1,3}([^_]+)_{1,3}', r'\1', content)

    return content.strip()


def extract_text_from_txt(content: bytes) -> str:
    """从纯文本提取"""
    # 尝试多种编码
    encodings = ['utf-8', 'gbk', 'gb2312', 'gb18030', 'latin-1']

    for encoding in encodings:
        try:
            return content.decode(encoding).strip()
        except (UnicodeDecodeError, AttributeError):
            continue

    # 最后尝试忽略错误
    return content.decode('utf-8', errors='ignore').strip()


def parse_document(
    file_content: bytes,
    file_extension: str,
    original_filename: str = ""
) -> str:
    """统一文档解析入口

    Args:
        file_content: 文件二进制内容
        file_extension: 文件扩展名（.pdf, .docx, .md, .txt）
        original_filename: 原始文件名

    Returns:
        提取的文本内容
    """
    ext = file_extension.lower().strip().lstrip('.')

    if ext == 'pdf':
        return extract_text_from_pdf(file_content)

    elif ext == 'docx':
        return extract_text_from_docx(file_content)
    elif ext == 'doc':
        return extract_text_from_doc(file_content)

    elif ext in ('md', 'markdown'):
        # Markdown 文件支持多编码
        if isinstance(file_content, bytes):
            # 尝试多种编码
            content = None
            for encoding in ['utf-8', 'gbk', 'gb2312', 'gb18030', 'latin-1']:
                try:
                    content = file_content.decode(encoding)
                    break
                except (UnicodeDecodeError, AttributeError):
                    continue
            if content is None:
                content = file_content.decode('utf-8', errors='ignore')
        else:
            content = file_content
        return extract_text_from_markdown(content)

    elif ext in ('txt', 'text'):
        return extract_text_from_txt(file_content)

    else:
        raise ValueError(f"不支持的文件格式: .{ext}，支持的格式: PDF, Word(.docx), Markdown(.md), TXT")


def truncate_for_analysis(text: str, max_chars: int = 50000) -> str:
    """截断文本以适应 AI 分析限制"""
    if len(text) <= max_chars:
        return text

    # 保留开头和结尾，中间部分截断
    head = text[:max_chars // 2]
    tail = text[-max_chars // 2:]

    return f"""{head}

[... 内容过长，中间部分已省略 ...]

{tail}"""
