"""Knowledge file import route: import JSON/Excel knowledge points."""

import io
import json
import logging
from typing import Any, Dict, List

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.knowledge import KnowledgePoint
from app.models.question import ExamCategory, ExamType
from app.models.user import User
from app.utils.security import require_teacher_or_admin

logger = logging.getLogger(__name__)
router = APIRouter(tags=["knowledge"])


@router.post("/import")
async def import_knowledge_points(
    file: UploadFile = File(...),
    parent_id: int | None = Form(None),
    category_id: int | None = Form(None),
    exam_type_id: int | None = Form(None),
    category: str | None = Form(None),
    exam_type: str | None = Form(None),
    current_user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
):
    """导入知识点文件（支持 JSON 和 Excel）

    - JSON 格式：数组或嵌套对象，支持 name/description/children/excerpt
    - Excel 格式：表头 name, description, parent_name（可选）
    """
    if not file.filename:
        raise HTTPException(status_code=400, detail="未上传文件")

    file_content = await file.read()
    await file.close()

    if not file_content:
        raise HTTPException(status_code=400, detail="文件内容为空")

    # 解析文件内容
    ext = file.filename.lower().split(".")[-1]
    nodes: List[Dict[str, Any]] = []

    if ext == "json":
        try:
            data = json.loads(file_content.decode("utf-8"))
        except json.JSONDecodeError as e:
            raise HTTPException(status_code=400, detail=f"JSON 解析失败: {e}") from e

        if isinstance(data, dict):
            data = [data]
        if not isinstance(data, list):
            raise HTTPException(status_code=400, detail="JSON 格式错误：应为数组或对象数组")

        def flatten(item: Dict[str, Any], parent_name: str | None = None) -> List[Dict[str, Any]]:
            result = []
            name = item.get("name") or item.get("名称") or item.get("title") or ""
            if not name:
                return result
            result.append(
                {
                    "name": str(name)[:100],
                    "description": (item.get("description") or item.get("描述") or "")[:2000],
                    "excerpt": (item.get("excerpt") or item.get("原文") or "")[:2000],
                    "parent_name": parent_name,
                }
            )
            children = item.get("children") or item.get("子节点") or []
            if isinstance(children, list):
                for child in children:
                    if isinstance(child, dict):
                        result.extend(flatten(child, name))
            return result

        for item in data:
            if isinstance(item, dict):
                nodes.extend(flatten(item))
    elif ext in ("xlsx", "xls"):
        try:
            import openpyxl
        except ImportError as e:
            raise HTTPException(status_code=500, detail="缺少 openpyxl 依赖") from e

        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Excel 解析失败: {e}") from e

        ws = wb.active
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(cell).lower().strip() if cell else "" for cell in row]

        name_idx = next((i for i, h in enumerate(headers) if h == "name"), None)
        desc_idx = next((i for i, h in enumerate(headers) if h == "description"), None)
        parent_idx = next((i for i, h in enumerate(headers) if h == "parent_name"), None)

        if name_idx is None:
            raise HTTPException(status_code=400, detail="Excel 缺少 name 列")

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            name = row[name_idx] if name_idx is not None and name_idx < len(row) else None
            if not name:
                continue
            nodes.append(
                {
                    "name": str(name)[:100],
                    "description": (str(row[desc_idx]) if desc_idx is not None and desc_idx < len(row) and row[desc_idx] else "")[:2000],
                    "excerpt": "",
                    "parent_name": str(row[parent_idx]) if parent_idx is not None and parent_idx < len(row) and row[parent_idx] else None,
                }
            )
        wb.close()
    else:
        raise HTTPException(status_code=400, detail=f"不支持的文件格式: {ext}，仅支持 json/xlsx")

    if not nodes:
        raise HTTPException(status_code=400, detail="文件中未找到有效知识点")

    # 解析 category/exam_type
    effective_category = category or "default"
    effective_exam_type = exam_type
    if category_id:
        cat_obj = db.query(ExamCategory).filter(ExamCategory.id == category_id).first()
        if cat_obj:
            effective_category = cat_obj.code
    if exam_type_id:
        et_obj = db.query(ExamType).filter(ExamType.id == exam_type_id).first()
        if et_obj:
            effective_exam_type = et_obj.code

    # 验证父节点
    effective_parent_id = parent_id
    if effective_parent_id:
        parent = db.query(KnowledgePoint).filter(KnowledgePoint.id == effective_parent_id).first()
        if not parent:
            raise HTTPException(status_code=404, detail="父知识点不存在")
        if parent.category == "default" or not parent.category:
            parent.category = effective_category
        if effective_exam_type:
            parent.exam_type = effective_exam_type

    # 构建 parent_name -> id 映射
    name_to_id = {}
    if effective_parent_id:
        parent = db.query(KnowledgePoint).filter(KnowledgePoint.id == effective_parent_id).first()
        if parent:
            name_to_id[parent.name] = effective_parent_id

    # 创建知识点
    created = []
    for node in nodes:
        parent_name = node.get("parent_name")
        resolved_parent_id = effective_parent_id
        if parent_name and parent_name in name_to_id:
            resolved_parent_id = name_to_id[parent_name]
        elif parent_name:
            # 创建父节点
            parent_kp = KnowledgePoint(
                name=parent_name[:100],
                parent_id=effective_parent_id,
                category=effective_category,
                exam_type=effective_exam_type,
                order=len(created),
                status=1,
                created_by=current_user.id,
            )
            db.add(parent_kp)
            db.flush()
            name_to_id[parent_name] = parent_kp.id
            resolved_parent_id = parent_kp.id
            created.append(parent_kp)

        kp = KnowledgePoint(
            name=node["name"],
            parent_id=resolved_parent_id,
            category=effective_category,
            exam_type=effective_exam_type,
            description=node.get("description") or None,
            content_excerpt=node.get("excerpt") or None,
            order=len(created),
            status=1,
            created_by=current_user.id,
        )
        db.add(kp)
        db.flush()
        created.append(kp)

    db.commit()
    for kp in created:
        db.refresh(kp)

    return {
        "success": True,
        "message": f"成功导入 {len(created)} 个知识点",
        "created_count": len(created),
        "items": [
            {
                "id": kp.id,
                "name": kp.name,
                "parent_id": kp.parent_id,
            }
            for kp in created
        ],
    }
