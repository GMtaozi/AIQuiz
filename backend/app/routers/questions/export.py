"""Questions Export Router - Export questions to Excel, Word, PDF"""

from datetime import datetime
import io
import logging
import os
from typing import List

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload

from app.database import get_db
from app.models.question import Question
from app.models.user import User
from app.routers.questions.utils import (
    _format_answer,
    _get_difficulty_name,
    _get_type_name,
    _safe_cell_value,
)
from app.utils.security import require_teacher_or_admin

logger = logging.getLogger(__name__)

router = APIRouter()


def _register_chinese_font() -> str | None:
    """评估 P2-16：中文字体路径配置化 + 多平台候选探测。

    原实现硬编码 Windows 路径 `C:/Windows/Fonts/simsun.ttc`，
    在 Linux 部署时导出 PDF 中文会回退 Helvetica（乱码/方块）。
    优先使用 settings.pdf_font_path，其次探测常见系统字体。
    返回注册后的字体内部名，找不到返回 None。
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    from app.config import settings

    candidates = []
    configured = getattr(settings, "pdf_font_path", "") or ""
    if configured:
        candidates.append(configured)
    candidates.extend(
        [
            "C:/Windows/Fonts/simsun.ttc",  # Windows 宋体
            "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",  # Linux 文泉驿微米黑
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",  # Linux Noto CJK
            "/System/Library/Fonts/PingFang.ttc",  # macOS 苹方
        ]
    )
    for path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("ChineseFont", path))
                return "ChineseFont"
            except Exception:
                continue
    return None


def _generate_excel(questions: List[Question]):
    """生成 Excel 文件"""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "题目列表"

    headers = ["序号", "题型", "难度", "题目内容", "选项A", "选项B", "选项C", "选项D", "正确答案", "解析"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, q in enumerate(questions, 2):
        ws.cell(row=row_idx, column=2, value=_get_type_name(q.question_type))
        ws.cell(row=row_idx, column=3, value=_get_difficulty_name(q.difficulty))
        ws.cell(row=row_idx, column=4, value=_safe_cell_value(q.content))
        ws.cell(row=row_idx, column=4).alignment = Alignment(wrap_text=True)

        options = sorted(q.options, key=lambda x: x.order) if q.options else []
        option_labels = ["A", "B", "C", "D"]
        for opt_idx, opt in enumerate(options[:4]):
            ws.cell(
                row=row_idx,
                column=5 + opt_idx,
                value=_safe_cell_value(f"{opt.option_label}. {opt.option_content}"),
            )

        ws.cell(row=row_idx, column=9, value=_safe_cell_value(_format_answer(q.answer, q.options, q.question_type)))
        ws.cell(row=row_idx, column=10, value=_safe_cell_value(q.explanation or ""))
        ws.cell(row=row_idx, column=10).alignment = Alignment(wrap_text=True)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 25
    ws.column_dimensions["G"].width = 25
    ws.column_dimensions["H"].width = 25
    ws.column_dimensions["I"].width = 10
    ws.column_dimensions["J"].width = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"questions_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    return output.getvalue(), filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _generate_word(questions: List[Question]):
    """生成 Word 文件"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    title = doc.add_heading("题目列表", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    stats_para = doc.add_paragraph()
    stats_para.add_run(f"总题数：{len(questions)}").bold = True
    stats_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph()

    questions_by_type = {}
    for q in questions:
        qtype = q.question_type
        if qtype not in questions_by_type:
            questions_by_type[qtype] = []
        questions_by_type[qtype].append(q)

    type_order = ["single_choice", "multiple_choice", "true_false", "essay"]
    type_names = {
        "single_choice": "一、单项选择题",
        "multiple_choice": "二、多项选择题",
        "true_false": "三、判断题",
        "essay": "四、简答题",
    }

    q_index = 1
    for qtype in type_order:
        if qtype not in questions_by_type:
            continue

        qs = questions_by_type[qtype]
        if not qs:
            continue

        section_title = doc.add_paragraph()
        section_title.add_run(type_names.get(qtype, qtype)).bold = True
        section_title.add_run(f"（共 {len(qs)} 题）")

        for q in qs:
            q_para = doc.add_paragraph()
            q_para.add_run(f"{q_index}. ").bold = True
            q_para.add_run(f"【{_get_type_name(q.question_type)}】")
            q_para.add_run(f"（{_get_difficulty_name(q.difficulty)}）")
            q_para.add_run(f" {q.content}")

            if q.question_type in ("single_choice", "multiple_choice"):
                options = sorted(q.options, key=lambda x: x.order) if q.options else []
                for opt in options:
                    opt_para = doc.add_paragraph(style="List Bullet")
                    opt_para.add_run(f"{opt.option_label}. {opt.option_content}")

            if q.answer:
                ans_para = doc.add_paragraph()
                ans_para.add_run(f"答案：{_format_answer(q.answer, q.options, q.question_type)}").italic = True

            if q.explanation:
                exp_para = doc.add_paragraph()
                exp_para.add_run(f"解析：{q.explanation}").italic = True

            doc.add_paragraph()
            q_index += 1

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)

    filename = f"questions_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx"
    return output.getvalue(), filename, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _generate_pdf(questions: List[Question]):
    """生成 PDF 文件"""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    chinese_font = _register_chinese_font() or "Helvetica"

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=styles["Heading1"],
        fontName=chinese_font,
        fontSize=16,
        alignment=TA_CENTER,
        spaceAfter=10,
    )

    info_style = ParagraphStyle(
        "ChineseInfo",
        parent=styles["Normal"],
        fontName=chinese_font,
        fontSize=10,
        alignment=TA_LEFT,
        spaceAfter=5,
    )

    section_style = ParagraphStyle(
        "ChineseSection",
        parent=styles["Normal"],
        fontName=chinese_font,
        fontSize=12,
        spaceBefore=10,
        spaceAfter=5,
    )

    question_style = ParagraphStyle(
        "ChineseQuestion",
        parent=styles["Normal"],
        fontName=chinese_font,
        fontSize=10,
        leading=14,
        spaceAfter=4,
    )

    option_style = ParagraphStyle(
        "ChineseOption",
        parent=styles["Normal"],
        fontName=chinese_font,
        fontSize=9,
        leading=12,
        leftIndent=15,
        spaceAfter=2,
    )

    answer_style = ParagraphStyle(
        "ChineseAnswer",
        parent=styles["Normal"],
        fontName=chinese_font,
        fontSize=9,
        textColor=colors.blue,
        leftIndent=15,
        spaceAfter=2,
    )

    story = []

    story.append(Paragraph("题目列表", title_style))
    story.append(Paragraph(f"总题数：{len(questions)}", info_style))
    story.append(Spacer(1, 15))

    questions_by_type = {}
    for q in questions:
        qtype = q.question_type
        if qtype not in questions_by_type:
            questions_by_type[qtype] = []
        questions_by_type[qtype].append(q)

    type_order = ["single_choice", "multiple_choice", "true_false", "essay"]
    type_names = {
        "single_choice": "一、单项选择题",
        "multiple_choice": "二、多项选择题",
        "true_false": "三、判断题",
        "essay": "四、简答题",
    }

    q_index = 1
    for qtype in type_order:
        if qtype not in questions_by_type:
            continue

        qs = questions_by_type[qtype]
        if not qs:
            continue

        story.append(Paragraph(f"{type_names.get(qtype, qtype)}（共 {len(qs)} 题）", section_style))

        for q in qs:
            q_text = (
                f"{q_index}. 【{_get_type_name(q.question_type)}】（{_get_difficulty_name(q.difficulty)}） {q.content}"
            )
            story.append(Paragraph(q_text, question_style))

            if q.question_type in ("single_choice", "multiple_choice"):
                options = sorted(q.options, key=lambda x: x.order) if q.options else []
                for opt in options:
                    story.append(Paragraph(f"{opt.option_label}. {opt.option_content}", option_style))

            if q.answer:
                story.append(Paragraph(f"答案：{_format_answer(q.answer, q.options, q.question_type)}", answer_style))

            if q.explanation:
                explanation_style = ParagraphStyle(
                    "ChineseExplanation",
                    parent=styles["Normal"],
                    fontName=chinese_font,
                    fontSize=9,
                    textColor=colors.grey,
                    leftIndent=15,
                    spaceAfter=2,
                )
                story.append(Paragraph(f"解析：{q.explanation}", explanation_style))

            story.append(Spacer(1, 5))
            q_index += 1

    doc.build(story)
    output.seek(0)

    filename = f"questions_{datetime.now().strftime('%Y%m%d%H%M%S')}.pdf"
    return output.getvalue(), filename, "application/pdf"


@router.get("/export")
def export_questions(
    subject_id: int | None = None,
    question_type: str | None = None,
    difficulty: int | None = None,
    question_ids: str | None = Query(None, description="逗号分隔的题目ID列表，如: 1,2,3"),
    format: str = Query("excel", pattern="^(excel|word|pdf)$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher_or_admin),
):
    """导出题目"""
    query = db.query(Question).options(joinedload(Question.options))

    if question_ids:
        ids = [int(x.strip()) for x in question_ids.split(",") if x.strip().isdigit()]
        if ids:
            query = query.filter(Question.id.in_(ids))
    else:
        if subject_id is not None:
            query = query.filter(Question.subject_id == subject_id)
        if question_type is not None:
            query = query.filter(Question.question_type == question_type)
        if difficulty is not None:
            query = query.filter(Question.difficulty == difficulty)

    questions = query.all()

    if not questions:
        raise HTTPException(status_code=404, detail="没有找到符合条件的题目")

    if format == "excel":
        file_content, filename, media_type = _generate_excel(questions)
    elif format == "word":
        file_content, filename, media_type = _generate_word(questions)
    else:
        file_content, filename, media_type = _generate_pdf(questions)

    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{filename}",
        "Access-Control-Allow-Origin": "*",
    }

    return StreamingResponse(iter([file_content]), media_type=media_type, headers=headers)
